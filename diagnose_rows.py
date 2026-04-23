"""
Диагностический скрипт для проверки конкретных строк
Проверяет почему не происходит мэппинг для определенных строк
"""

import sys
import openpyxl
from collections import defaultdict


# Проблемные строки
DEBUG_TEMPLATE_ROWS = [2335, 2857]
DEBUG_PROVODKI_ROWS = [2194, 2203]


def normalize_value(value, is_ofr=False):
    """Нормализация значения как в оригинальном скрипте"""
    if value is None or str(value).strip() == "":
        return "" if is_ofr else None
    
    # Преобразование с учетом типа
    if isinstance(value, (int, float)):
        if isinstance(value, float):
            if value == int(value):
                value_str = str(int(value))
            else:
                value_str = f"{value:.10f}".rstrip('0').rstrip('.')
        else:
            value_str = str(value)
    else:
        value_str = str(value).strip()
    
    # Удаление точек
    normalized = value_str.replace(".", "")
    
    if not normalized.isdigit() and normalized != "":
        return None
    
    return normalized


def parse_provodki_value(value):
    """Парсинг значения из проводок"""
    if value is None or str(value).strip() == "":
        return None
    
    # Преобразование с учетом типа
    if isinstance(value, (int, float)):
        if isinstance(value, float):
            if value == int(value):
                value_str = str(int(value))
            else:
                value_str = f"{value:.10f}".rstrip('0').rstrip('.')
        else:
            value_str = str(value)
    else:
        value_str = str(value).strip()
    
    # Подчеркивания
    underscore_count = value_str.count("_")
    
    if underscore_count > 1:
        return None
    
    if underscore_count == 1:
        parts = value_str.split("_")
        account = parts[0]
        ofr = parts[1]
    else:
        account = value_str
        ofr = ""
    
    # Удаление точек
    account_norm = account.replace(".", "")
    ofr_norm = ofr.replace(".", "")
    
    if not account_norm.isdigit():
        return None
    
    if ofr_norm and not ofr_norm.isdigit():
        return None
    
    return (account_norm, ofr_norm)


def create_key(dt_account, dt_ofr, kt_account, kt_ofr):
    """Создание ключа"""
    return f"{dt_account}{dt_ofr}{kt_account}{kt_ofr}"


def diagnose_template_rows(wb, rows):
    """Диагностика строк шаблона"""
    print("\n" + "=" * 80)
    print("ДИАГНОСТИКА СТРОК ШАБЛОНА")
    print("=" * 80)
    
    sheet_name = "Шаблон пров и атриб"
    ws = wb[sheet_name]
    header_row = 7
    
    # Получаем индексы столбцов
    headers = {}
    for idx, cell in enumerate(ws[header_row], start=1):
        if cell.value:
            headers[str(cell.value).strip()] = idx
    
    col_filter = headers.get("Признак плана счетов (ПАО/КИБ)")
    col_dt_account = headers.get("Номер счета по Дт")
    col_dt_ofr = headers.get("Символ ОФР Дт")
    col_kt_account = headers.get("Номер счета по Кт")
    col_kt_ofr = headers.get("Символ ОФР Кт")
    col_pokaza = headers.get("№ Показа")
    
    keys = {}
    
    for row_num in rows:
        print(f"\n--- Строка {row_num} ---")
        
        # Фильтр
        filter_val = ws.cell(row=row_num, column=col_filter).value
        print(f"Признак плана счетов: '{filter_val}'")
        
        if filter_val is None or str(filter_val).strip() != "ПАО":
            print("⚠️ ПРОПУЩЕНО: не ПАО")
            continue
        
        # Читаем значения
        dt_account_val = ws.cell(row=row_num, column=col_dt_account).value
        dt_ofr_val = ws.cell(row=row_num, column=col_dt_ofr).value
        kt_account_val = ws.cell(row=row_num, column=col_kt_account).value
        kt_ofr_val = ws.cell(row=row_num, column=col_kt_ofr).value
        pokaza_val = ws.cell(row=row_num, column=col_pokaza).value
        
        print(f"\nИсходные значения:")
        print(f"  Дт счет:     type={type(dt_account_val).__name__:10s} value='{dt_account_val}' repr={repr(dt_account_val)}")
        print(f"  Дт ОФР:      type={type(dt_ofr_val).__name__:10s} value='{dt_ofr_val}' repr={repr(dt_ofr_val)}")
        print(f"  Кт счет:     type={type(kt_account_val).__name__:10s} value='{kt_account_val}' repr={repr(kt_account_val)}")
        print(f"  Кт ОФР:      type={type(kt_ofr_val).__name__:10s} value='{kt_ofr_val}' repr={repr(kt_ofr_val)}")
        print(f"  № Показа:    type={type(pokaza_val).__name__:10s} value='{pokaza_val}'")
        
        # Нормализация
        dt_account_norm = normalize_value(dt_account_val, False)
        dt_ofr_norm = normalize_value(dt_ofr_val, True)
        kt_account_norm = normalize_value(kt_account_val, False)
        kt_ofr_norm = normalize_value(kt_ofr_val, True)
        
        if dt_ofr_norm is None:
            dt_ofr_norm = ""
        if kt_ofr_norm is None:
            kt_ofr_norm = ""
        
        print(f"\nНормализованные значения:")
        print(f"  Дт счет:     '{dt_account_norm}' [длина={len(dt_account_norm) if dt_account_norm else 0}]")
        print(f"  Дт ОФР:      '{dt_ofr_norm}' [длина={len(dt_ofr_norm)}]")
        print(f"  Кт счет:     '{kt_account_norm}' [длина={len(kt_account_norm) if kt_account_norm else 0}]")
        print(f"  Кт ОФР:      '{kt_ofr_norm}' [длина={len(kt_ofr_norm)}]")
        
        if dt_account_norm is None or kt_account_norm is None:
            print("❌ ОШИБКА: невалидные счета")
            continue
        
        # Ключ
        key = create_key(dt_account_norm, dt_ofr_norm, kt_account_norm, kt_ofr_norm)
        print(f"\nКлюч сравнения:")
        print(f"  '{key}' [длина={len(key)}]")
        print(f"  Байты: {key.encode('utf-8')}")
        
        keys[row_num] = key
    
    return keys


def diagnose_provodki_rows(wb, rows):
    """Диагностика строк проводок"""
    print("\n" + "=" * 80)
    print("ДИАГНОСТИКА СТРОК ПРОВОДОК")
    print("=" * 80)
    
    sheet_name = "Проводки по СП анализ"
    ws = wb[sheet_name]
    header_row = 1
    
    # Получаем индексы столбцов
    headers = {}
    for idx, cell in enumerate(ws[header_row], start=1):
        if cell.value:
            headers[str(cell.value).strip()] = idx
    
    col_dt = headers.get("Дт")
    col_kt = headers.get("Кт")
    col_expanded = headers.get("Номера строк развёрнутых проводок")
    
    keys = {}
    
    for row_num in rows:
        print(f"\n--- Строка {row_num} ---")
        
        # Проверка развёрнутых
        expanded_val = ws.cell(row=row_num, column=col_expanded).value
        print(f"Развёрнутые проводки: '{expanded_val}'")
        
        if expanded_val and str(expanded_val).strip():
            print("⚠️ ПРОПУЩЕНО: есть развёрнутые проводки (обрабатывается отдельно)")
            continue
        
        # Читаем значения
        dt_val = ws.cell(row=row_num, column=col_dt).value
        kt_val = ws.cell(row=row_num, column=col_kt).value
        
        print(f"\nИсходные значения:")
        print(f"  Дт: type={type(dt_val).__name__:10s} value='{dt_val}' repr={repr(dt_val)}")
        print(f"  Кт: type={type(kt_val).__name__:10s} value='{kt_val}' repr={repr(kt_val)}")
        
        # Парсинг
        dt_parsed = parse_provodki_value(dt_val)
        kt_parsed = parse_provodki_value(kt_val)
        
        if dt_parsed is None or kt_parsed is None:
            print("❌ ОШИБКА: не удалось распарсить")
            continue
        
        dt_account_norm, dt_ofr_norm = dt_parsed
        kt_account_norm, kt_ofr_norm = kt_parsed
        
        print(f"\nРаспарсенные значения:")
        print(f"  Дт счет:     '{dt_account_norm}' [длина={len(dt_account_norm)}]")
        print(f"  Дт ОФР:      '{dt_ofr_norm}' [длина={len(dt_ofr_norm)}]")
        print(f"  Кт счет:     '{kt_account_norm}' [длина={len(kt_account_norm)}]")
        print(f"  Кт ОФР:      '{kt_ofr_norm}' [длина={len(kt_ofr_norm)}]")
        
        # Ключ
        key = create_key(dt_account_norm, dt_ofr_norm, kt_account_norm, kt_ofr_norm)
        print(f"\nКлюч сравнения:")
        print(f"  '{key}' [длина={len(key)}]")
        print(f"  Байты: {key.encode('utf-8')}")
        
        keys[row_num] = key
    
    return keys


def compare_keys(template_keys, provodki_keys):
    """Сравнение ключей"""
    print("\n" + "=" * 80)
    print("СРАВНЕНИЕ КЛЮЧЕЙ")
    print("=" * 80)
    
    print(f"\nКлючи из шаблона:")
    for row, key in template_keys.items():
        print(f"  Строка {row}: '{key}'")
    
    print(f"\nКлючи из проводок:")
    for row, key in provodki_keys.items():
        print(f"  Строка {row}: '{key}'")
    
    print(f"\nПроверка совпадений:")
    template_set = set(template_keys.values())
    provodki_set = set(provodki_keys.values())
    
    matches = template_set & provodki_set
    
    if matches:
        print(f"✅ Найдено совпадений: {len(matches)}")
        for key in matches:
            template_rows = [r for r, k in template_keys.items() if k == key]
            provodki_rows = [r for r, k in provodki_keys.items() if k == key]
            print(f"  Ключ '{key}':")
            print(f"    Шаблон строки: {template_rows}")
            print(f"    Проводки строки: {provodki_rows}")
    else:
        print(f"❌ Совпадений НЕ найдено!")
        print(f"\nВозможные причины:")
        print(f"  1. Разные типы данных (float vs string)")
        print(f"  2. Проблемы точности float")
        print(f"  3. Невидимые символы в ячейках")
        print(f"  4. Разная обработка пустых символов ОФР")


def main():
    """Главная функция"""
    print("=" * 80)
    print("ДИАГНОСТИЧЕСКИЙ СКРИПТ ДЛЯ ПРОВЕРКИ ПРОБЛЕМНЫХ СТРОК")
    print("=" * 80)
    
    # Получение файла
    if len(sys.argv) > 1:
        input_file = sys.argv[1]
    else:
        input_file = input("Введите путь к Excel-файлу: ").strip()
    
    print(f"\nОткрываем файл: {input_file}")
    
    # Открываем с data_only=True (как в оригинале)
    wb = openpyxl.load_workbook(input_file, data_only=True)
    
    # Диагностика
    template_keys = diagnose_template_rows(wb, DEBUG_TEMPLATE_ROWS)
    provodki_keys = diagnose_provodki_rows(wb, DEBUG_PROVODKI_ROWS)
    
    # Сравнение
    compare_keys(template_keys, provodki_keys)
    
    wb.close()
    
    print("\n" + "=" * 80)
    print("ДИАГНОСТИКА ЗАВЕРШЕНА")
    print("=" * 80)
    
    return 0


if __name__ == "__main__":
    sys.exit(main())
