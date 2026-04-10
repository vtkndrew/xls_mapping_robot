"""
Скрипт для автоматического мэппинга проводок между листами Excel
Автор: Mapping Robot
Версия: 1.0
"""

import sys
import openpyxl
from openpyxl.utils import get_column_letter
from datetime import datetime
from collections import defaultdict
import re
import os


class ProcessLogger:
    """Класс для логирования ошибок и предупреждений"""
    
    def __init__(self, log_file="errors.txt"):
        self.log_file = log_file
        self.errors = []
        self.warnings = []
    
    def log_error(self, sheet, row, column, value, reason):
        """Записать ошибку в лог"""
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        message = f"[{timestamp}] ОШИБКА | Лист: {sheet} | Строка: {row} | Столбец: {column} | Значение: {value} | Причина: {reason}"
        self.errors.append(message)
    
    def log_warning(self, sheet, row, column, value, reason):
        """Записать предупреждение в лог"""
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        message = f"[{timestamp}] ПРЕДУПРЕЖДЕНИЕ | Лист: {sheet} | Строка: {row} | Столбец: {column} | Значение: {value} | Причина: {reason}"
        self.warnings.append(message)
    
    def save(self):
        """Сохранить логи в файл"""
        with open(self.log_file, 'w', encoding='utf-8') as f:
            f.write("=== ЛОГ ОБРАБОТКИ ===\n")
            f.write(f"Дата и время: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n")
            
            if self.errors:
                f.write("ОШИБКИ:\n")
                for error in self.errors:
                    f.write(error + "\n")
                f.write("\n")
            else:
                f.write("ОШИБКИ: Нет\n\n")
            
            if self.warnings:
                f.write("ПРЕДУПРЕЖДЕНИЯ:\n")
                for warning in self.warnings:
                    f.write(warning + "\n")
            else:
                f.write("ПРЕДУПРЕЖДЕНИЯ: Нет\n")


class ProcessStatistics:
    """Класс для сбора статистики обработки"""
    
    def __init__(self):
        # Статистика по шаблону
        self.template_total = 0
        self.template_processed = 0
        self.template_errors = 0
        self.template_matched = 0
        
        # Статистика по проводкам
        self.provodki_total = 0
        self.provodki_processed = 0
        self.provodki_expanded = 0
        self.provodki_errors = 0
        self.provodki_matched = 0
    
    def save(self, filename="statistics.txt", output_file=""):
        """Сохранить статистику в файл"""
        with open(filename, 'w', encoding='utf-8') as f:
            f.write("=== СТАТИСТИКА ОБРАБОТКИ ===\n")
            f.write(f"Дата и время: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n")
            
            f.write('Лист "Шаблон пров и атриб":\n')
            f.write(f"  Всего строк: {self.template_total}\n")
            f.write(f"  Обработано строк: {self.template_processed}\n")
            f.write(f"  Успешно сматчено строк Шаблон → Проводки: {self.template_matched}\n")
            f.write(f"  Ошибок в шаблоне: {self.template_errors}\n\n")
            
            f.write('Лист "Проводки по СП анализ":\n')
            f.write(f"  Всего строк: {self.provodki_total}\n")
            f.write(f"  Обработано строк: {self.provodki_processed}\n")
            f.write(f"  Обработано развёрнутых проводок: {self.provodki_expanded}\n")
            f.write(f"  Успешно сматчено строк Проводки → Шаблон: {self.provodki_matched}\n")
            f.write(f"  Ошибок в проводках: {self.provodki_errors}\n\n")
            
            total_errors = self.template_errors + self.provodki_errors
            f.write(f"Общее количество ошибок: {total_errors}\n\n")
            
            if output_file:
                f.write(f"Результат сохранён в: {output_file}\n")


def get_input_file():
    """Получить путь к входному Excel-файлу"""
    if len(sys.argv) > 1:
        input_file = sys.argv[1]
    else:
        input_file = input("Введите путь к Excel-файлу: ").strip()
    
    if not input_file:
        raise ValueError("Ошибка: путь не указан.")
    
    if not os.path.exists(input_file):
        raise FileNotFoundError(f"Ошибка: файл '{input_file}' не найден.")
    
    return input_file


def validate_file_structure(wb):
    """Проверить наличие необходимых листов и столбцов"""
    required_sheets = {
        "Шаблон пров и атриб": {
            "header_row": 7,
            "columns": [
                "Признак плана счетов (ПАО/КИБ)",
                "Номер счета по Дт",
                "Символ ОФР Дт",
                "Номер счета по Кт",
                "Символ ОФР Кт",
                "№ Показа",
                "Дт Мэпинг Робот",
                "Кт Мэпинг Робот"
            ]
        },
        "Проводки по СП анализ": {
            "header_row": 1,
            "columns": [
                "Дт",
                "Кт",
                "Номера строк развёрнутых проводок",
                "Дт Мэпинг Робот",
                "Кт Мэпинг Робот",
                "СП Робот"
            ]
        }
    }
    
    # Проверка наличия листов
    for sheet_name in required_sheets.keys():
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Ошибка: лист '{sheet_name}' не найден в файле.")
    
    # Проверка наличия столбцов
    for sheet_name, config in required_sheets.items():
        ws = wb[sheet_name]
        header_row = config["header_row"]
        headers = []
        
        # Читаем заголовки
        for cell in ws[header_row]:
            if cell.value:
                headers.append(str(cell.value).strip())
        
        # Проверяем наличие всех необходимых столбцов
        for column in config["columns"]:
            if column not in headers:
                raise ValueError(f"Ошибка: столбец '{column}' не найден на листе '{sheet_name}'.")
    
    print("✓ Валидация структуры файла пройдена успешно")


def get_column_index_by_name(ws, column_name, header_row):
    """Получить индекс столбца по его имени"""
    for idx, cell in enumerate(ws[header_row], start=1):
        if cell.value and str(cell.value).strip() == column_name:
            return idx
    return None


def normalize_template_value(value, row_num, column_name, sheet_name, logger):
    """
    Нормализовать значение для листа Шаблон:
    - Удалить точки
    - Проверить, что остались только цифры
    - Пустое значение допустимо только для символов ОФР
    """
    # Обработка None или пустого значения
    if value is None or str(value).strip() == "":
        # Пустое значение допустимо только для символов ОФР
        if "Символ ОФР" in column_name:
            return ""
        else:
            logger.log_error(sheet_name, row_num, column_name, value, 
                           "Пустое значение недопустимо для счета")
            return None
    
    # Преобразуем в строку
    value_str = str(value).strip()
    
    # Удаляем точки
    normalized = value_str.replace(".", "")
    
    # Проверяем, что остались только цифры
    if not normalized.isdigit() and normalized != "":
        logger.log_error(sheet_name, row_num, column_name, value_str, 
                       "Недопустимые символы (разрешены только цифры и точки)")
        return None
    
    return normalized


def parse_provodki_account(value, row_num, column_name, sheet_name, logger):
    """
    Парсить счет из листа Проводки:
    - Формат: счет_символОФР или просто счет
    - Разделить по подчеркиванию
    - Удалить точки
    - Проверить на допустимые символы
    - Возвращает (счет, символОФР) или None
    """
    # Обработка None или пустого значения
    if value is None or str(value).strip() == "":
        logger.log_error(sheet_name, row_num, column_name, value, 
                       "Пустое значение недопустимо для счета")
        return None
    
    # Преобразуем в строку
    value_str = str(value).strip()
    
    # Подсчитываем количество подчеркиваний
    underscore_count = value_str.count("_")
    
    if underscore_count > 1:
        logger.log_error(sheet_name, row_num, column_name, value_str, 
                       "Более одного подчеркивания в значении")
        return None
    
    # Разделяем по подчеркиванию
    if underscore_count == 1:
        parts = value_str.split("_")
        account = parts[0]
        ofr = parts[1]
    else:
        account = value_str
        ofr = ""
    
    # Удаляем точки из обеих частей
    account_norm = account.replace(".", "")
    ofr_norm = ofr.replace(".", "")
    
    # Проверяем, что в счете только цифры
    if not account_norm.isdigit():
        logger.log_error(sheet_name, row_num, column_name, value_str, 
                       f"Недопустимые символы в счете '{account}'")
        return None
    
    # Проверяем, что в символе ОФР только цифры (если не пустой)
    if ofr_norm and not ofr_norm.isdigit():
        logger.log_error(sheet_name, row_num, column_name, value_str, 
                       f"Недопустимые символы в символе ОФР '{ofr}'")
        return None
    
    return (account_norm, ofr_norm)


def create_comparison_key(dt_account, dt_ofr, kt_account, kt_ofr):
    """Создать ключ сравнения из 4 компонентов"""
    return f"{dt_account}{dt_ofr}{kt_account}{kt_ofr}"


def create_excel_reference(sheet_name, column_letter, row_number):
    """Создать ссылку Excel формата ='Лист'!A1"""
    return f"='{sheet_name}'!{column_letter}{row_number}"


def join_references(references_list):
    """Объединить список ссылок через запятую с пробелом"""
    return ", ".join(references_list)


def process_template_sheet(wb_read, logger, stats):
    """
    Обработать лист Шаблон (Этап 1)
    Возвращает словарь ключей сравнения: {key: [список строк]}
    """
    sheet_name = "Шаблон пров и атриб"
    ws = wb_read[sheet_name]
    header_row = 7
    data_start_row = 8
    
    # Получаем индексы столбцов
    col_filter = get_column_index_by_name(ws, "Признак плана счетов (ПАО/КИБ)", header_row)
    col_dt_account = get_column_index_by_name(ws, "Номер счета по Дт", header_row)
    col_dt_ofr = get_column_index_by_name(ws, "Символ ОФР Дт", header_row)
    col_kt_account = get_column_index_by_name(ws, "Номер счета по Кт", header_row)
    col_kt_ofr = get_column_index_by_name(ws, "Символ ОФР Кт", header_row)
    col_pokazа = get_column_index_by_name(ws, "№ Показа", header_row)
    
    # Словарь для хранения ключей: {comparison_key: [{'row': row_num, 'pokaza': value, 'dt_col': col, 'kt_col': col}]}
    template_keys = defaultdict(list)
    
    processed_count = 0
    error_count = 0
    
    # Обрабатываем строки
    max_row = ws.max_row
    stats.template_total = max_row - data_start_row + 1
    
    for row_num in range(data_start_row, max_row + 1):
        # Фильтрация по ПАО
        filter_value = ws.cell(row=row_num, column=col_filter).value
        if filter_value is None or str(filter_value).strip() != "ПАО":
            continue
        
        # Читаем значения
        dt_account_val = ws.cell(row=row_num, column=col_dt_account).value
        dt_ofr_val = ws.cell(row=row_num, column=col_dt_ofr).value
        kt_account_val = ws.cell(row=row_num, column=col_kt_account).value
        kt_ofr_val = ws.cell(row=row_num, column=col_kt_ofr).value
        pokaza_val = ws.cell(row=row_num, column=col_pokazа).value
        
        # Нормализуем значения
        dt_account_norm = normalize_template_value(dt_account_val, row_num, "Номер счета по Дт", sheet_name, logger)
        dt_ofr_norm = normalize_template_value(dt_ofr_val, row_num, "Символ ОФР Дт", sheet_name, logger)
        kt_account_norm = normalize_template_value(kt_account_val, row_num, "Номер счета по Кт", sheet_name, logger)
        kt_ofr_norm = normalize_template_value(kt_ofr_val, row_num, "Символ ОФР Кт", sheet_name, logger)
        
        # Если хотя бы одно значение счета невалидно - пропускаем строку
        if dt_account_norm is None or kt_account_norm is None:
            error_count += 1
            continue
        
        # Символы ОФР могут быть None (заменяем на пустую строку)
        if dt_ofr_norm is None:
            dt_ofr_norm = ""
        if kt_ofr_norm is None:
            kt_ofr_norm = ""
        
        # Создаём ключ сравнения
        comparison_key = create_comparison_key(dt_account_norm, dt_ofr_norm, kt_account_norm, kt_ofr_norm)
        
        # Сохраняем информацию о строке
        template_keys[comparison_key].append({
            'row': row_num,
            'pokaza': pokaza_val if pokaza_val else "",
            'dt_col': get_column_letter(col_dt_account),
            'kt_col': get_column_letter(col_kt_account)
        })
        
        processed_count += 1
    
    stats.template_processed = processed_count
    stats.template_errors = error_count
    
    print(f"✓ Обработано строк шаблона (фильтр ПАО): {processed_count}")
    print(f"  Уникальных ключей: {len(template_keys)}")
    print(f"  Ошибок валидации: {error_count}")
    
    return template_keys


def process_provodki_sheet_stage1(wb_read, template_keys, logger, stats):
    """
    Обработать лист Проводки (Этап 2 - первая часть)
    Создать словарь ключей для проводок без развёрнутых
    Возвращает словарь ключей, список результатов для записи и mapping_cache
    """
    sheet_name = "Проводки по СП анализ"
    ws = wb_read[sheet_name]
    header_row = 1
    data_start_row = 2
    
    # Получаем индексы столбцов
    col_dt = get_column_index_by_name(ws, "Дт", header_row)
    col_kt = get_column_index_by_name(ws, "Кт", header_row)
    col_expanded = get_column_index_by_name(ws, "Номера строк развёрнутых проводок", header_row)
    col_dt_mapping = get_column_index_by_name(ws, "Дт Мэпинг Робот", header_row)
    col_kt_mapping = get_column_index_by_name(ws, "Кт Мэпинг Робот", header_row)
    col_sp_robot = get_column_index_by_name(ws, "СП Робот", header_row)
    
    # Словари для хранения данных
    provodki_keys = defaultdict(list)
    results_template = []  # Результаты для записи в Шаблон
    results_provodki = []  # Результаты для записи в Проводки
    mapping_cache = {}  # Кеш ссылок для этапа 3 (развёрнутые проводки)
    
    processed_count = 0
    error_count = 0
    
    # Обрабатываем строки
    max_row = ws.max_row
    stats.provodki_total = max_row - data_start_row + 1
    
    for row_num in range(data_start_row, max_row + 1):
        # Проверяем, есть ли развёрнутые проводки
        expanded_val = ws.cell(row=row_num, column=col_expanded).value
        if expanded_val and str(expanded_val).strip():
            # Эти строки обрабатываются на этапе 3
            continue
        
        # Читаем значения
        dt_val = ws.cell(row=row_num, column=col_dt).value
        kt_val = ws.cell(row=row_num, column=col_kt).value
        
        # Парсим счета
        dt_parsed = parse_provodki_account(dt_val, row_num, "Дт", sheet_name, logger)
        kt_parsed = parse_provodki_account(kt_val, row_num, "Кт", sheet_name, logger)
        
        # Если парсинг не удался - пропускаем
        if dt_parsed is None or kt_parsed is None:
            error_count += 1
            continue
        
        dt_account_norm, dt_ofr_norm = dt_parsed
        kt_account_norm, kt_ofr_norm = kt_parsed
        
        # Создаём ключ сравнения
        comparison_key = create_comparison_key(dt_account_norm, dt_ofr_norm, kt_account_norm, kt_ofr_norm)
        
        # Сохраняем информацию о строке
        provodki_keys[comparison_key].append({
            'row': row_num,
            'dt_col': get_column_letter(col_dt),
            'kt_col': get_column_letter(col_kt)
        })
        
        processed_count += 1
    
    stats.provodki_processed = processed_count
    stats.provodki_errors = error_count
    
    print(f"✓ Обработано строк проводок (без развёрнутых): {processed_count}")
    print(f"  Уникальных ключей: {len(provodki_keys)}")
    print(f"  Ошибок валидации: {error_count}")
    
    # Теперь сопоставляем ключи
    # Этап 1: Шаблон → Проводки (запись в Шаблон)
    template_matched = 0
    for key, template_rows in template_keys.items():
        if key in provodki_keys:
            # Нашли совпадение
            provodki_rows = provodki_keys[key]
            
            # Для каждой строки шаблона
            for template_info in template_rows:
                # Собираем ссылки на Дт и Кт из проводок
                dt_refs = []
                kt_refs = []
                for provodki_info in provodki_rows:
                    dt_ref = create_excel_reference("Проводки по СП анализ", 
                                                   provodki_info['dt_col'], 
                                                   provodki_info['row'])
                    kt_ref = create_excel_reference("Проводки по СП анализ", 
                                                   provodki_info['kt_col'], 
                                                   provodki_info['row'])
                    dt_refs.append(dt_ref)
                    kt_refs.append(kt_ref)
                
                # Записываем в результаты для Шаблона
                results_template.append({
                    'sheet': 'Шаблон пров и атриб',
                    'row': template_info['row'],
                    'dt_mapping_col': get_column_index_by_name(wb_read['Шаблон пров и атриб'], 
                                                              "Дт Мэпинг Робот", 7),
                    'kt_mapping_col': get_column_index_by_name(wb_read['Шаблон пров и атриб'], 
                                                              "Кт Мэпинг Робот", 7),
                    'dt_value': join_references(dt_refs),
                    'kt_value': join_references(kt_refs)
                })
                template_matched += 1
    
    stats.template_matched = template_matched
    print(f"✓ Сматчено строк Шаблон → Проводки: {template_matched}")
    
    # Этап 2: Проводки → Шаблон (запись в Проводки)
    provodki_matched = 0
    for key, provodki_rows in provodki_keys.items():
        if key in template_keys:
            # Нашли совпадение
            template_rows = template_keys[key]
            
            # Для каждой строки проводок
            for provodki_info in provodki_rows:
                # Собираем ссылки на Номер счета по Дт и Кт из шаблона
                dt_refs = []
                kt_refs = []
                pokaza_values = []
                
                for template_info in template_rows:
                    dt_ref = create_excel_reference("Шаблон пров и атриб", 
                                                   template_info['dt_col'], 
                                                   template_info['row'])
                    kt_ref = create_excel_reference("Шаблон пров и атриб", 
                                                   template_info['kt_col'], 
                                                   template_info['row'])
                    dt_refs.append(dt_ref)
                    kt_refs.append(kt_ref)
                    
                    # Собираем уникальные значения "№ Показа"
                    pokaza = str(template_info['pokaza']).strip() if template_info['pokaza'] else ""
                    if pokaza and pokaza not in pokaza_values:
                        pokaza_values.append(pokaza)
                
                # Формируем значения ссылок
                dt_value = join_references(dt_refs)
                kt_value = join_references(kt_refs)
                
                # Записываем в результаты для Проводок
                results_provodki.append({
                    'sheet': 'Проводки по СП анализ',
                    'row': provodki_info['row'],
                    'dt_mapping_col': col_dt_mapping,
                    'kt_mapping_col': col_kt_mapping,
                    'sp_robot_col': col_sp_robot,
                    'dt_value': dt_value,
                    'kt_value': kt_value,
                    'sp_value': ", ".join(pokaza_values)
                })
                
                # Сохраняем ссылки в кеш для использования на этапе 3
                mapping_cache[provodki_info['row']] = {
                    'dt_value': dt_value,
                    'kt_value': kt_value
                }
                
                provodki_matched += 1
    
    stats.provodki_matched = provodki_matched
    print(f"✓ Сматчено строк Проводки → Шаблон: {provodki_matched}")
    print(f"✓ Ссылок сохранено в кеш для этапа 3: {len(mapping_cache)}")
    
    return provodki_keys, results_template, results_provodki, mapping_cache


def process_expanded_provodki(wb_read, mapping_cache, logger, stats):
    """
    Обработать развёрнутые проводки (Этап 3)
    Использует mapping_cache для получения ссылок из этапа 2
    Возвращает список результатов для записи
    """
    sheet_name = "Проводки по СП анализ"
    ws = wb_read[sheet_name]
    header_row = 1
    data_start_row = 2
    
    # Получаем индексы столбцов
    col_expanded = get_column_index_by_name(ws, "Номера строк развёрнутых проводок", header_row)
    col_dt_mapping = get_column_index_by_name(ws, "Дт Мэпинг Робот", header_row)
    col_kt_mapping = get_column_index_by_name(ws, "Кт Мэпинг Робот", header_row)
    
    results = []
    expanded_count = 0
    
    # Обрабатываем строки
    max_row = ws.max_row
    
    for row_num in range(data_start_row, max_row + 1):
        # Проверяем, есть ли развёрнутые проводки
        expanded_val = ws.cell(row=row_num, column=col_expanded).value
        if not expanded_val or not str(expanded_val).strip():
            continue
        
        # Парсим номера строк
        expanded_str = str(expanded_val).strip()
        # Разделяем по запятой
        row_numbers = []
        for part in expanded_str.split(','):
            part = part.strip()
            if part.isdigit():
                row_numbers.append(int(part))
        
        if not row_numbers:
            logger.log_warning(sheet_name, row_num, "Номера строк развёрнутых проводок", 
                             expanded_str, "Не удалось распарсить номера строк")
            continue
        
        # Собираем ссылки из указанных строк (используя кеш из этапа 2)
        dt_refs = []
        kt_refs = []
        
        for ref_row in row_numbers:
            # Читаем значения из кеша (сформированные на этапе 2)
            if ref_row in mapping_cache:
                cached_data = mapping_cache[ref_row]
                
                # Добавляем Дт ссылки
                if cached_data.get('dt_value'):
                    dt_refs.append(cached_data['dt_value'])
                
                # Добавляем Кт ссылки
                if cached_data.get('kt_value'):
                    kt_refs.append(cached_data['kt_value'])
            else:
                # Строка не была обработана на этапе 2 - логируем предупреждение
                logger.log_warning(sheet_name, row_num, "Номера строк развёрнутых проводок",
                                 ref_row, f"Строка {ref_row} не найдена в кеше мэппинга")
        
        # Записываем результаты
        if dt_refs or kt_refs:
            results.append({
                'sheet': sheet_name,
                'row': row_num,
                'dt_mapping_col': col_dt_mapping,
                'kt_mapping_col': col_kt_mapping,
                'dt_value': join_references(dt_refs) if dt_refs else "",
                'kt_value': join_references(kt_refs) if kt_refs else ""
            })
            expanded_count += 1
    
    stats.provodki_expanded = expanded_count
    print(f"✓ Обработано развёрнутых проводок: {expanded_count}")
    
    return results


def write_results_to_excel(input_file, all_results):
    """
    Записать результаты в Excel файл
    """
    # Создаём имя выходного файла
    base_name = os.path.splitext(input_file)[0]
    output_file = f"{base_name}_processed.xlsx"
    
    # Открываем файл для записи (без data_only, чтобы сохранить формулы)
    wb_write = openpyxl.load_workbook(input_file, data_only=False)
    
    # Записываем результаты
    for result in all_results:
        sheet_name = result['sheet']
        ws = wb_write[sheet_name]
        row = result['row']
        
        # Записываем Дт Мэпинг Робот
        if 'dt_value' in result and result['dt_value']:
            ws.cell(row=row, column=result['dt_mapping_col']).value = result['dt_value']
        
        # Записываем Кт Мэпинг Робот
        if 'kt_value' in result and result['kt_value']:
            ws.cell(row=row, column=result['kt_mapping_col']).value = result['kt_value']
        
        # Записываем СП Робот (если есть)
        if 'sp_value' in result and result['sp_value']:
            ws.cell(row=row, column=result['sp_robot_col']).value = result['sp_value']
    
    # Сохраняем файл
    wb_write.save(output_file)
    wb_write.close()
    
    print(f"✓ Результаты записаны в файл: {output_file}")
    
    return output_file


def main():
    """Главная функция"""
    try:
        print("=" * 60)
        print("СКРИПТ АВТОМАТИЧЕСКОГО МЭППИНГА ПРОВОДОК")
        print("=" * 60)
        print()
        
        # 1. Получение пути к файлу
        print("Шаг 1: Получение пути к файлу...")
        input_file = get_input_file()
        print(f"✓ Файл: {input_file}")
        print()
        
        # 2. Инициализация логгера и статистики
        logger = ProcessLogger("errors.txt")
        stats = ProcessStatistics()
        
        # 3. Открытие файла для чтения значений
        print("Шаг 2: Открытие файла...")
        wb_read = openpyxl.load_workbook(input_file, data_only=True)
        print("✓ Файл открыт для чтения")
        print()
        
        # 4. Валидация структуры
        print("Шаг 3: Валидация структуры файла...")
        validate_file_structure(wb_read)
        print()
        
        # 5. Обработка Шаблона (Этап 1)
        print("Шаг 4: Обработка листа 'Шаблон пров и атриб'...")
        template_keys = process_template_sheet(wb_read, logger, stats)
        print()
        
        # 6. Обработка Проводок (Этап 2)
        print("Шаг 5: Обработка листа 'Проводки по СП анализ'...")
        provodki_keys, results_template, results_provodki, mapping_cache = process_provodki_sheet_stage1(
            wb_read, template_keys, logger, stats)
        print()
        
        # 7. Обработка развёрнутых проводок (Этап 3)
        print("Шаг 6: Обработка развёрнутых проводок...")
        results_expanded = process_expanded_provodki(wb_read, mapping_cache, logger, stats)
        print()
        
        # Закрываем workbook для чтения
        wb_read.close()
        
        # 8. Запись результатов
        print("Шаг 7: Запись результатов...")
        all_results = results_template + results_provodki + results_expanded
        output_file = write_results_to_excel(input_file, all_results)
        print()
        
        # 9. Сохранение логов и статистики
        print("Шаг 8: Сохранение логов и статистики...")
        logger.save()
        stats.save("statistics.txt", output_file)
        print("✓ Логи сохранены в: errors.txt")
        print("✓ Статистика сохранена в: statistics.txt")
        print()
        
        print("=" * 60)
        print("ОБРАБОТКА ЗАВЕРШЕНА УСПЕШНО!")
        print("=" * 60)
        print(f"Результат: {output_file}")
        print(f"Логи ошибок: errors.txt")
        print(f"Статистика: statistics.txt")
        print()
        
        return 0
        
    except Exception as e:
        print()
        print("=" * 60)
        print("КРИТИЧЕСКАЯ ОШИБКА!")
        print("=" * 60)
        print(f"Ошибка: {e}")
        print()
        import traceback
        traceback.print_exc()
        return 1


if __name__ == "__main__":
    sys.exit(main())
